"""Route common study-document formats into numbered text units."""

from pathlib import Path
import os
import shutil
import subprocess
import tempfile

from ingestion.pdf_loader import load_pdf_pages
from ingestion.pptx_loader import load_pptx_slides


DOCUMENT_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".docx", ".txt", ".md"}


def load_document(path: str | Path) -> list[tuple[int, str]]:
    path = Path(path)
    extension = path.suffix.lower()
    if extension == ".pdf":
        return load_pdf_pages(str(path))
    if extension == ".pptx":
        return load_pptx_slides(str(path))
    if extension == ".docx":
        from docx import Document
        document = Document(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            parts.extend(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
        return [(1, "\n".join(parts))] if parts else []
    if extension in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace").strip()
        return [(1, text)] if text else []
    if extension == ".ppt":
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice and os.name == "nt":
            candidate = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
            if candidate.is_file():
                soffice = str(candidate)
        if not soffice:
            raise ValueError("Legacy .ppt needs LibreOffice installed for conversion to .pptx")
        with tempfile.TemporaryDirectory() as directory:
            profile = (Path(directory) / "profile").as_uri()
            result = subprocess.run(
                [soffice, f"-env:UserInstallation={profile}", "--headless",
                 "--convert-to", "pptx", "--outdir", directory, str(path)],
                capture_output=True, text=True, timeout=120, check=False,
            )
            converted = Path(directory) / (path.stem + ".pptx")
            if result.returncode != 0 or not converted.exists():
                raise ValueError(f"Could not convert {path.name}: {result.stderr.strip()[:200]}")
            return load_pptx_slides(str(converted))
    raise ValueError(f"Unsupported document format: {extension or '(no extension)'}")
