"""
ingestion/pptx_loader.py

Extracts text from a PowerPoint file, slide by slide, including both
the visible slide text and any speaker notes (notes often contain the
actual explanation, which is valuable for RAG).
"""

from typing import List, Tuple

from pptx import Presentation


def load_pptx_slides(file_path: str) -> List[Tuple[int, str]]:
    """
    Extracts text from a PPTX file, one entry per slide.

    Returns:
        A list of (slide_number, combined_text) tuples, 1-indexed.
        combined_text includes visible text boxes + speaker notes.
        Blank slides are skipped.
    """
    presentation = Presentation(file_path)
    slides_out = []

    for i, slide in enumerate(presentation.slides, start=1):
        text_parts = []

        # Visible text on the slide (titles, bullet points, text boxes)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        text_parts.append(line.strip())

        # Speaker notes, if present
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text
            if notes_text and notes_text.strip():
                text_parts.append(f"[Notes: {notes_text.strip()}]")

        combined = "\n".join(text_parts)
        if combined.strip():
            slides_out.append((i, combined.strip()))

    return slides_out


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python pptx_loader.py <path_to_pptx>")
        sys.exit(1)

    result = load_pptx_slides(sys.argv[1])
    print(f"Extracted {len(result)} slides with text.")
    if result:
        first_slide_num, first_slide_text = result[0]
        preview = first_slide_text[:200].replace("\n", " ")
        print(f"Slide {first_slide_num} preview: {preview}...")
